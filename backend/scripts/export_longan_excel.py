"""从源 Excel 完整导出龙眼未标注记录（含原表全部列与文件说明）。"""

import scripts._bootstrap  # noqa: F401

from config import CACHE_FILE, OUTPUT_DIR, ensure_output_dir

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from food_inspection.parser import _clean, _find_header_row, _read_all_sheets

CACHE = Path(CACHE_FILE)
OUT = Path(OUTPUT_DIR) / "龙眼未标注完整导出.xlsx"

META_COLUMNS = ["导出序号", "数据来源文件", "数据来源路径", "工作表", "公告标题", "公告说明"]
NOTE = (
    "说明：以下数据直接从源 Excel 原表抽取，保留全部原始列。"
    "源文件中无「不合格项目名称」「实测值」「标准值」等检测明细列，"
    "故系统无法自动标注具体不合格项目。"
)


def _load_target_keys() -> list[dict[str, str]]:
    with CACHE.open(encoding="utf-8") as f:
        uq = json.load(f)["unqualified"]
    return [
        {
            "source_file": r.get("source_file") or "",
            "product": r.get("product") or "",
            "company": r.get("company") or "",
        }
        for r in uq
        if "龙眼" in (r.get("product") or "")
        and not (r.get("unqualified_item") or "").strip()
    ]


def _collect_preamble(rows: list[list[Any]], header_idx: int) -> tuple[str, str]:
    title_parts: list[str] = []
    note_parts: list[str] = []
    for idx in range(header_idx):
        text = _clean(" ".join(_clean(c) for c in rows[idx] if _clean(c)))
        if not text:
            continue
        if idx == 0 or "附件" in text or len(text) < 40:
            title_parts.append(text)
        else:
            note_parts.append(text)
    return " / ".join(title_parts), " ".join(note_parts)


def _row_to_dict(headers: list[str], cells: list[Any]) -> dict[str, str]:
    row: dict[str, str] = {}
    for idx, header in enumerate(headers):
        name = header or f"列{idx + 1}"
        if name in row:
            name = f"{name}_{idx + 1}"
        value = _clean(cells[idx]) if idx < len(cells) else ""
        row[name] = value
    return row


def _extract_full_rows(targets: list[dict[str, str]]) -> tuple[list[dict[str, str]], list[str]]:
    source_files = sorted({t["source_file"] for t in targets if t["source_file"]})
    all_rows: list[dict[str, str]] = []
    column_order: list[str] = []

    for source_file in source_files:
        path = Path(source_file)
        if not path.exists():
            continue
        for sheet_name, rows, _ in _read_all_sheets(source_file):
            header_idx = _find_header_row(rows)
            if header_idx is None:
                continue
            headers = [_clean(c) or f"列{i + 1}" for i, c in enumerate(rows[header_idx])]
            title, preamble = _collect_preamble(rows, header_idx)

            for col in headers:
                if col not in column_order:
                    column_order.append(col)

            for cells in rows[header_idx + 1 :]:
                first = _clean(cells[0]) if cells else ""
                if not first:
                    continue
                if first in ("序号", "备注"):
                    continue
                if first.startswith("（声明") or first.startswith("声明"):
                    break
                if not first.replace(".", "").isdigit():
                    if all_rows:
                        break
                    continue

                row_dict = _row_to_dict(headers, list(cells))
                product = row_dict.get("食品名称") or row_dict.get("产品名称") or row_dict.get("样品名称") or ""
                if "龙眼" not in product:
                    continue

                full_row = {
                    "数据来源文件": path.name,
                    "数据来源路径": str(path),
                    "工作表": sheet_name,
                    "公告标题": title,
                    "公告说明": preamble,
                }
                full_row.update(row_dict)
                all_rows.append(full_row)

    final_columns = column_order
    return all_rows, final_columns


def _style_header(ws, row: int, col_count: int) -> None:
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(bold=True, color="FFFFFF", size=11)
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="B4C6E7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = align
        cell.border = border


def _style_body(ws, start_row: int, end_row: int, col_count: int) -> None:
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    alt_fill = PatternFill("solid", fgColor="F2F7FB")
    for row in range(start_row, end_row + 1):
        for col in range(1, col_count + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if row % 2 == 0:
                cell.fill = alt_fill


def _autosize_columns(ws, headers: list[str], data_rows: list[dict[str, str]], start_row: int) -> None:
    for col_idx, header in enumerate(headers, 1):
        letter = get_column_letter(col_idx)
        max_len = len(header)
        for row in data_rows:
            value = str(row.get(header, ""))
            max_len = max(max_len, min(len(value), 50))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 48)


def export() -> Path:
    targets = _load_target_keys()
    rows, body_columns = _extract_full_rows(targets)
    headers = META_COLUMNS + body_columns

    wb = Workbook()

    ws = wb.active
    ws.title = "龙眼未标注明细"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = ws.cell(row=1, column=1, value=f"龙眼未标注记录完整导出（共 {len(rows)} 条）")
    title_cell.font = Font(bold=True, size=14, color="1F4E79")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    note_cell = ws.cell(row=2, column=1, value=NOTE)
    note_cell.font = Font(size=10, color="666666")
    note_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 40

    header_row = 3
    for col, name in enumerate(headers, 1):
        ws.cell(row=header_row, column=col, value=name)
    _style_header(ws, header_row, len(headers))

    data_start = header_row + 1
    export_rows: list[dict[str, str]] = []
    for i, row in enumerate(rows, 1):
        export_row = {"导出序号": str(i)}
        export_row.update(row)
        export_rows.append(export_row)
        excel_row = data_start + i - 1
        for col, key in enumerate(headers, 1):
            ws.cell(row=excel_row, column=col, value=export_row.get(key, ""))

    if export_rows:
        _style_body(ws, data_start, data_start + len(export_rows) - 1, len(headers))
    _autosize_columns(ws, headers, export_rows, data_start)
    ws.freeze_panes = "A4"
    if export_rows:
        ws.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{data_start + len(export_rows) - 1}"

    info = wb.create_sheet("字段说明")
    info["A1"] = "原始 Excel 列说明"
    info["A1"].font = Font(bold=True, size=12)
    info_rows = [
        ("导出序号", "导出文件中的顺序编号"),
        ("数据来源文件", "源 Excel 文件名"),
        ("数据来源路径", "源 Excel 完整路径"),
        ("工作表", "源 Excel 工作表名称"),
        ("公告标题", "表头前的标题/附件说明"),
        ("公告说明", "表头前的抽检说明文字"),
    ]
    for col in body_columns:
        if col in ("序号", "被抽样单位名称", "被抽样单位地址", "标称生产企业名称", "标称生产企业地址"):
            info_rows.append((col, "源表原始列"))
        elif col == "食品名称":
            info_rows.append((col, "源表原始列（产品名称）"))
        elif col == "型号规格":
            info_rows.append((col, "源表原始列（规格）"))
        elif col == "生产日期/批号":
            info_rows.append((col, "源表原始列"))
        elif col == "承检机构":
            info_rows.append((col, "源表原始列（检测机构）"))
        elif col == "备注":
            info_rows.append((col, "源表原始列"))
        else:
            info_rows.append((col, "源表原始列"))

    info["A2"] = "字段名"
    info["B2"] = "说明"
    _style_header(info, 2, 2)
    for idx, (name, desc) in enumerate(info_rows, 3):
        info.cell(row=idx, column=1, value=name)
        info.cell(row=idx, column=2, value=desc)
    info.column_dimensions["A"].width = 28
    info.column_dimensions["B"].width = 60

    ensure_output_dir()
    wb.save(OUT)
    return OUT


if __name__ == "__main__":
    print(export())
