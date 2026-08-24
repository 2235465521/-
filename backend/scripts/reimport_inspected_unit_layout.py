#!/usr/bin/env python3
"""重导「受检单位 + 抽样单位」表头结构的 Excel（修正抽样单位误作被抽检单位）。"""

from __future__ import annotations

import argparse
import os
import sys
from pathlib import Path

BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from openpyxl import load_workbook  # noqa: E402

from food_inspection.db import connection, resolve_disk_path  # noqa: E402
from food_inspection.mysql_v2 import create_compat_view, insert_records_batch  # noqa: E402
from food_inspection.parser import iter_excel_files, parse_file  # noqa: E402
from food_inspection.parser.fields import source_relative_key  # noqa: E402
from food_inspection.parser.text import _clean, _normalize_header  # noqa: E402

_BATCH = 500


def _sheet_has_inspected_sampling_layout(path: str) -> bool:
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return False
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows(min_row=1, max_row=15, values_only=True):
                headers = [_normalize_header(_clean(c)) for c in row if _clean(c)]
                if not headers:
                    continue
                has_inspected = any(
                    "受检单位" in h and "地址" not in h for h in headers
                )
                has_sampling = any(
                    h in ("抽样单位", "抽样单位名称", "抽检单位", "抽检单位名称")
                    or (h.startswith("抽样单位") and not h.startswith("被抽样"))
                    for h in headers
                )
                if has_inspected and has_sampling:
                    return True
        return False
    finally:
        wb.close()


def _collect_layout_files() -> list[str]:
    matched: list[str] = []
    for fp in iter_excel_files():
        if _sheet_has_inspected_sampling_layout(fp):
            rel = source_relative_key(fp)
            if rel:
                matched.append(rel)
    return sorted(set(matched))


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description="重导受检单位+抽样单位表头 Excel")
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    rel_keys = _collect_layout_files()
    if not rel_keys:
        print("未找到匹配的源文件")
        return 0

    print(f"匹配源文件 {len(rel_keys)} 个")
    for rel in rel_keys:
        print(f"  {rel}")

    deleted = imported = qualified = unqualified = 0
    missing = 0

    with connection() as conn, conn.cursor() as cur:
        if not args.dry_run:
            for rel in rel_keys:
                cur.execute("DELETE FROM inspection_base WHERE file_source = %s", (rel,))
                deleted += cur.rowcount

            cur.execute("SELECT COALESCE(MAX(id), 0) + 1 AS n FROM inspection_base")
            next_id = int(cur.fetchone()["n"])
        else:
            cur.execute(
                "SELECT COUNT(*) AS n FROM inspection_base WHERE file_source IN %s",
                (tuple(rel_keys),),
            )
            deleted = int(cur.fetchone()["n"])
            next_id = 0

        pending: list = []
        for rel in rel_keys:
            path = resolve_disk_path(rel)
            if not path or not os.path.isfile(path):
                missing += 1
                print(f"跳过（文件不存在）: {rel}")
                continue
            records = parse_file(path)
            for record in records:
                if args.dry_run:
                    imported += 1
                    if record.status == "qualified":
                        qualified += 1
                    else:
                        unqualified += 1
                    continue
                pending.append((record, next_id))
                next_id += 1
                imported += 1
                if record.status == "qualified":
                    qualified += 1
                else:
                    unqualified += 1
                if len(pending) >= _BATCH:
                    insert_records_batch(cur, pending)
                    pending.clear()

        if not args.dry_run and pending:
            insert_records_batch(cur, pending)
            conn.commit()
            create_compat_view()

    print(f"删除旧记录: {deleted}")
    print(f"重新导入: {imported}（合格 {qualified}，不合格 {unqualified}）")
    if missing:
        print(f"缺失源文件: {missing}")
    if args.dry_run:
        print("（dry-run，未写入数据库）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
