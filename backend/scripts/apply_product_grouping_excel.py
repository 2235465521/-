#!/usr/bin/env python3
"""根据「产品归类核查表.xlsx」中人工校正结果，批量更新数据库 food_name。

默认 dry-run；加 --apply 才真正写库。
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

import scripts._bootstrap  # noqa: F401

from openpyxl import load_workbook

from food_inspection.db import connection


def _load_mapping(excel_path: Path, sheet_name: str = "归类明细") -> dict[str, str]:
    wb = load_workbook(excel_path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise SystemExit(f"未找到工作表：{sheet_name}")
    ws = wb[sheet_name]
    rows = ws.iter_rows(values_only=True)
    headers = [str(h or "").strip() for h in next(rows)]
    try:
        idx_original = headers.index("原始产品名")
        idx_corrected_sub = headers.index("校正后小类")
        idx_proposed_sub = headers.index("建议小类")
    except ValueError:
        # 兼容旧版
        idx_original = headers.index("原始产品名")
        idx_corrected_sub = headers.index("校正后归类") if "校正后归类" in headers else -1
        idx_proposed_sub = headers.index("建议归类") if "建议归类" in headers else -1

    mapping: dict[str, str] = {}
    for row in rows:
        if not row or idx_original >= len(row):
            continue
        original = str(row[idx_original] or "").strip()
        if not original:
            continue
        corrected = (
            str(row[idx_corrected_sub] or "").strip()
            if idx_corrected_sub >= 0 and idx_corrected_sub < len(row)
            else ""
        )
        proposed = (
            str(row[idx_proposed_sub] or "").strip()
            if idx_proposed_sub >= 0 and idx_proposed_sub < len(row)
            else ""
        )
        target = corrected or proposed or original
        if target != original:
            mapping[original] = target
    wb.close()
    return mapping


def apply_mapping(mapping: dict[str, str], *, dry_run: bool) -> tuple[int, int]:
    if not mapping:
        return 0, 0

    qualified_updates = 0
    unqualified_updates = 0
    with connection() as conn, conn.cursor() as cur:
        for original, target in mapping.items():
            if dry_run:
                cur.execute(
                    "SELECT COUNT(*) AS n FROM inspection_qualified WHERE food_name = %s",
                    (original,),
                )
                qualified_updates += int(cur.fetchone()["n"])
                cur.execute(
                    "SELECT COUNT(*) AS n FROM inspection_unqualified WHERE food_name = %s",
                    (original,),
                )
                unqualified_updates += int(cur.fetchone()["n"])
                continue

            cur.execute(
                "UPDATE inspection_qualified SET food_name = %s WHERE food_name = %s",
                (target, original),
            )
            qualified_updates += cur.rowcount
            cur.execute(
                "UPDATE inspection_unqualified SET food_name = %s WHERE food_name = %s",
                (target, original),
            )
            unqualified_updates += cur.rowcount
        if not dry_run:
            conn.commit()
    return qualified_updates, unqualified_updates


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description="应用产品归类 Excel 校正结果")
    parser.add_argument("excel", help="产品归类核查表.xlsx 路径")
    parser.add_argument("--apply", action="store_true", help="真正写库（默认仅预览）")
    args = parser.parse_args()

    mapping = _load_mapping(Path(args.excel))
    print(f"待归并映射：{len(mapping):,} 条（原始名 → 目标名）")
    q, u = apply_mapping(mapping, dry_run=not args.apply)
    mode = "预览" if not args.apply else "已更新"
    print(f"{mode}：合格表 {q:,} 行，不合格表 {u:,} 行")
    if not args.apply:
        print("确认无误后请加 --apply 执行写库")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
