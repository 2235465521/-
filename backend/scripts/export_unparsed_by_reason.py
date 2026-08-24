"""按原因分类导出仍未解析的 Excel 源文件路径。"""

from __future__ import annotations

import scripts._bootstrap  # noqa: F401

from config import OUTPUT_DIR, ensure_output_dir

import os
import sys
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from food_inspection.parser import (
    _build_column_map,
    _clean,
    _detect_file_mode,
    _detect_sheet_status,
    _extract_location_from_path,
    _find_header_row,
    _get_cell,
    _is_data_row,
    _read_all_sheets,
    _resolve_company,
    iter_excel_files,
    parse_file,
)

OUT = os.path.join(OUTPUT_DIR, "未解析文件-按原因分类.xlsx")

REASON_LABELS: dict[str, str] = {
    "read_failed_or_empty": "01-无法读取或空白",
    "exception": "02-读取异常",
    "unknown_status": "03-无法识别合格不合格",
    "no_header_row": "04-找不到表头",
    "missing_key_columns": "05-缺少产品或单位列",
    "no_data_rows": "06-有表头无数据行",
    "all_rows_no_company_product": "07-公司名产品名皆空",
    "rows_filtered": "08-数据行被过滤",
    "empty_sheet": "09-空工作表",
    "file_missing": "10-文件不存在",
    "unknown": "11-其他",
}

REASON_DESC: dict[str, str] = {
    "read_failed_or_empty": "Excel 打不开、损坏或所有工作表为空",
    "exception": "读取过程中抛出异常",
    "unknown_status": "不是合格/不合格明细表（汇总页、说明页、你点我检汇总等）",
    "no_header_row": "找不到含「序号」+ 产品名称的表头行",
    "missing_key_columns": "有表头但缺少食品名称/样品名称/产品名称，且无企业/单位列",
    "no_data_rows": "表头正确但下面没有有效数据行",
    "all_rows_no_company_product": "有数据行但每行公司名和产品名都为空",
    "rows_filtered": "有数据行但不符合当前行识别规则（需人工看样例）",
    "empty_sheet": "工作表没有任何内容",
    "file_missing": "文件路径不存在",
    "unknown": "未能归类的原因",
}

COLUMNS = [
    "序号",
    "未解析原因",
    "源文件路径",
    "源文件名",
    "文件夹省份",
    "文件夹城市",
    "工作表",
    "合格/不合格判定",
    "表头行号",
    "表头列名",
    "已识别列",
    "第1条数据样例",
]


def _diagnose_sheet(
    filepath: str,
    sheet_name: str,
    rows: list[list],
    sheet_text: str,
    file_mode: str,
) -> dict[str, str] | None:
    path_province, path_city = _extract_location_from_path(filepath)
    base = {
        "源文件路径": filepath,
        "源文件名": os.path.basename(filepath),
        "文件夹省份": path_province,
        "文件夹城市": path_city,
        "工作表": sheet_name,
        "合格/不合格判定": "",
        "表头行号": "",
        "表头列名": "",
        "已识别列": "",
        "第1条数据样例": "",
    }

    if not rows:
        return {**base, "未解析原因": REASON_LABELS["empty_sheet"]}

    status = _detect_sheet_status(sheet_name, sheet_text, file_mode)
    if status is None:
        return {**base, "未解析原因": REASON_LABELS["unknown_status"]}

    base["合格/不合格判定"] = "合格" if status == "qualified" else "不合格"
    header_idx = _find_header_row(rows)
    if header_idx is None:
        return {**base, "未解析原因": REASON_LABELS["no_header_row"]}

    headers = [_clean(c) for c in rows[header_idx]]
    col_map = _build_column_map(headers)
    base["表头行号"] = str(header_idx + 1)
    base["表头列名"] = " | ".join(h for h in headers if h)
    mapped = []
    for field, idx in sorted(col_map.items(), key=lambda x: x[1]):
        name = headers[idx] if idx < len(headers) else field
        mapped.append(f"{field}:{name}")
    base["已识别列"] = "；".join(mapped) if mapped else "(无)"

    if "product" not in col_map and "company" not in col_map and "sampled_unit" not in col_map:
        return {**base, "未解析原因": REASON_LABELS["missing_key_columns"]}

    data_rows = 0
    skipped_no_name = 0
    sample = ""
    for row in rows[header_idx + 1 :]:
        cells = list(row)
        first = _clean(cells[0]) if cells else ""
        if not first or first in ("序号", "备注"):
            continue
        if first.startswith("（声明") or first.startswith("声明"):
            break
        if not _is_data_row(cells, headers):
            if data_rows:
                break
            continue
        data_rows += 1
        if not sample:
            sample = " | ".join(_clean(c) for c in cells[: min(len(cells), 14)] if _clean(c))
        company = _resolve_company(cells, col_map)
        product = _get_cell(cells, col_map, "product")
        if not company and not product:
            skipped_no_name += 1

    base["第1条数据样例"] = sample

    if data_rows == 0:
        return {**base, "未解析原因": REASON_LABELS["no_data_rows"]}
    if skipped_no_name == data_rows:
        return {**base, "未解析原因": REASON_LABELS["all_rows_no_company_product"]}
    return {**base, "未解析原因": REASON_LABELS["rows_filtered"]}


def _diagnose_file(filepath: str) -> list[dict[str, str]]:
    if not os.path.isfile(filepath):
        return [
            {
                "未解析原因": REASON_LABELS["file_missing"],
                "源文件路径": filepath,
                "源文件名": os.path.basename(filepath),
                "文件夹省份": "",
                "文件夹城市": "",
                "工作表": "",
                "合格/不合格判定": "",
                "表头行号": "",
                "表头列名": "",
                "已识别列": "",
                "第1条数据样例": "",
            }
        ]

    try:
        if parse_file(filepath):
            return []
    except Exception:
        return [
            {
                "未解析原因": REASON_LABELS["exception"],
                "源文件路径": filepath,
                "源文件名": os.path.basename(filepath),
                "文件夹省份": _extract_location_from_path(filepath)[0],
                "文件夹城市": _extract_location_from_path(filepath)[1],
                "工作表": "",
                "合格/不合格判定": "",
                "表头行号": "",
                "表头列名": "",
                "已识别列": "",
                "第1条数据样例": "",
            }
        ]

    file_mode = _detect_file_mode(os.path.basename(filepath))
    sheets = _read_all_sheets(filepath)
    if not sheets:
        return [
            {
                "未解析原因": REASON_LABELS["read_failed_or_empty"],
                "源文件路径": filepath,
                "源文件名": os.path.basename(filepath),
                "文件夹省份": _extract_location_from_path(filepath)[0],
                "文件夹城市": _extract_location_from_path(filepath)[1],
                "工作表": "",
                "合格/不合格判定": "",
                "表头行号": "",
                "表头列名": "",
                "已识别列": "",
                "第1条数据样例": "",
            }
        ]

    issues: list[dict[str, str]] = []
    for sheet_name, rows, sheet_text in sheets:
        detail = _diagnose_sheet(filepath, sheet_name, rows, sheet_text, file_mode)
        if detail:
            issues.append(detail)
    return issues


def _style_header(ws, row: int, cols: int) -> None:
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(bold=True, color="FFFFFF", size=11)
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def export() -> str:
    sys.stdout.reconfigure(encoding="utf-8")
    files = list(iter_excel_files())
    by_reason: dict[str, list[dict[str, str]]] = defaultdict(list)
    ok_files = 0

    for i, filepath in enumerate(files, 1):
        issues = _diagnose_file(filepath)
        if not issues:
            ok_files += 1
        else:
            for item in issues:
                by_reason[item["未解析原因"]].append(item)
        if i % 100 == 0:
            print(f"已扫描 {i}/{len(files)}，成功 {ok_files}，未解析条目 {sum(len(v) for v in by_reason.values())}...", flush=True)

    wb = Workbook()
    summary = wb.active
    summary.title = "汇总"
    summary["A1"] = "未解析原因"
    summary["B1"] = "条数(工作表/文件级)"
    summary["C1"] = "说明"
    _style_header(summary, 1, 3)

    total_issues = sum(len(v) for v in by_reason.values())
    summary["A2"] = "解析成功(文件)"
    summary["B2"] = ok_files
    summary["C2"] = f"共 {len(files)} 个 Excel 文件"
    summary["A3"] = "未解析(条目合计)"
    summary["B3"] = total_issues
    summary["C3"] = "一个文件多个工作表可能有多条"

    row = 5
    for reason_key in sorted(REASON_LABELS.values()):
        items = by_reason.get(reason_key, [])
        if not items:
            continue
        code = reason_key.split("-", 1)[0]
        desc_key = [k for k, v in REASON_LABELS.items() if v == reason_key][0]
        summary.cell(row=row, column=1, value=reason_key)
        summary.cell(row=row, column=2, value=len(items))
        summary.cell(row=row, column=3, value=REASON_DESC.get(desc_key, ""))
        row += 1

    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 14
    summary.column_dimensions["C"].width = 56

    for reason_key in sorted(REASON_LABELS.values()):
        items = by_reason.get(reason_key, [])
        if not items:
            continue
        sheet_name = reason_key[:31]
        ws = wb.create_sheet(sheet_name)
        for col, name in enumerate(COLUMNS, 1):
            ws.cell(row=1, column=col, value=name)
        _style_header(ws, 1, len(COLUMNS))
        for idx, item in enumerate(items, 1):
            ws.cell(row=idx + 1, column=1, value=idx)
            for col, key in enumerate(COLUMNS[1:], 2):
                ws.cell(row=idx + 1, column=col, value=item.get(key, ""))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(items) + 1}"
        for col_idx, header in enumerate(COLUMNS, 1):
            letter = get_column_letter(col_idx)
            ws.column_dimensions[letter].width = 48 if header in ("源文件路径", "表头列名", "第1条数据样例") else 16

    ensure_output_dir()
    wb.save(OUT)
    print(f"\n总计文件: {len(files)}")
    print(f"解析成功: {ok_files}")
    print(f"未解析条目: {total_issues}")
    print(f"已导出: {OUT}")
    for reason_key in sorted(REASON_LABELS.values()):
        n = len(by_reason.get(reason_key, []))
        if n:
            print(f"  {reason_key}: {n}")
    return OUT


if __name__ == "__main__":
    export()
