"""导出未解析文件中占比最高的一类：表头有「序号」但首列不是「序号」。"""

from __future__ import annotations

import scripts._bootstrap  # noqa: F401

from config import OUTPUT_DIR, ensure_output_dir

import os
import sys
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from food_inspection.parser import (
    _build_column_map,
    _clean,
    _detect_file_mode,
    _detect_sheet_status,
    _extract_location_from_path,
    _find_header_row,
    _read_all_sheets,
    iter_excel_files,
    parse_file,
)

OUT = os.path.join(OUTPUT_DIR, "未解析-表头格式问题.xlsx")

COLUMNS = [
    "序号",
    "未解析原因",
    "源文件路径",
    "源文件名",
    "文件夹省份",
    "文件夹城市",
    "工作表",
    "合格/不合格",
    "表头行号(1起)",
    "表头第1列",
    "表头全部列名",
    "系统能否识别产品列",
    "系统能否识别单位列",
    "第1条数据样例",
]


def _find_header_candidate(rows: list[list]) -> tuple[int | None, str]:
    """返回 (行号, 原因)。与 _find_header_row 不同，用于诊断。"""
    for idx in range(min(25, len(rows))):
        cells = [_clean(c) for c in rows[idx]]
        if not cells:
            continue
        joined = "".join(cells)
        has_seq = "序号" in cells
        has_product_header = any(k in joined for k in ("食品名称", "样品名称", "产品名称"))
        if not has_seq or not has_product_header:
            continue
        if cells[0] == "序号":
            return idx, "标准表头(序号在第1列)"
        return idx, "序号不在第1列(前面有抽样编号等)"
    return None, ""


def _collect_rows(filepath: str) -> list[dict[str, str]]:
    if not os.path.isfile(filepath):
        return []

    try:
        if parse_file(filepath):
            return []
    except Exception:
        pass

    filename = os.path.basename(filepath)
    file_mode = _detect_file_mode(filename)
    path_province, path_city = _extract_location_from_path(filepath)
    out: list[dict[str, str]] = []

    for sheet_name, rows, sheet_text in _read_all_sheets(filepath):
        if not rows:
            continue
        status = _detect_sheet_status(sheet_name, sheet_text, file_mode)
        if status is None:
            continue
        if _find_header_row(rows) is not None:
            continue

        header_idx, reason = _find_header_candidate(rows)
        if header_idx is None:
            continue
        if reason == "标准表头(序号在第1列)":
            continue

        headers = [_clean(c) for c in rows[header_idx]]
        col_map = _build_column_map(headers)
        sample = ""
        for row in rows[header_idx + 1 : header_idx + 10]:
            cells = [_clean(c) for c in row]
            if not cells:
                continue
            marker = cells[1] if len(cells) > 1 and cells[1].replace(".", "").isdigit() else cells[0]
            if marker.replace(".", "").isdigit():
                sample = " | ".join(c for c in cells[: min(len(cells), 14)] if c)
                break

        out.append(
            {
                "未解析原因": reason,
                "源文件路径": filepath,
                "源文件名": filename,
                "文件夹省份": path_province,
                "文件夹城市": path_city,
                "工作表": sheet_name,
                "合格/不合格": "合格" if status == "qualified" else "不合格",
                "表头行号(1起)": str(header_idx + 1),
                "表头第1列": headers[0] if headers else "",
                "表头全部列名": " | ".join(h for h in headers if h),
                "系统能否识别产品列": "是" if "product" in col_map else "否",
                "系统能否识别单位列": "是" if ("company" in col_map or "sampled_unit" in col_map) else "否",
                "第1条数据样例": sample,
            }
        )

    return out


def _style_header(ws, row: int, col_count: int) -> None:
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(bold=True, color="FFFFFF", size=11)
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def export() -> str:
    sys.stdout.reconfigure(encoding="utf-8")
    files = list(iter_excel_files())
    all_rows: list[dict[str, str]] = []

    for i, filepath in enumerate(files, 1):
        all_rows.extend(_collect_rows(filepath))
        if i % 100 == 0:
            print(f"已扫描 {i}/{len(files)}，已收集 {len(all_rows)} 条...", flush=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "表头格式问题"

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))
    ws.cell(
        row=1,
        column=1,
        value=f"未解析最多的一类：表头「序号」不在第1列（共 {len(all_rows)} 个工作表）",
    ).font = Font(bold=True, size=14, color="1F4E79")

    for col, name in enumerate(COLUMNS, 1):
        ws.cell(row=2, column=col, value=name)
    _style_header(ws, 2, len(COLUMNS))

    for i, row in enumerate(all_rows, 1):
        ws.cell(row=2 + i, column=1, value=i)
        for col, key in enumerate(COLUMNS[1:], 2):
            ws.cell(row=2 + i, column=col, value=row.get(key, ""))

    if all_rows:
        ws.auto_filter.ref = f"A2:{get_column_letter(len(COLUMNS))}{2 + len(all_rows)}"
    ws.freeze_panes = "A3"

    for i, header in enumerate(COLUMNS, 1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = 48 if header in ("源文件路径", "表头全部列名", "第1条数据样例") else 16

    stats = wb.create_sheet("表头第1列统计")
    c1 = Counter(r["表头第1列"] for r in all_rows)
    stats["A1"], stats["B1"] = "表头第1列", "次数"
    _style_header(stats, 1, 2)
    for idx, (name, count) in enumerate(c1.most_common(), 2):
        stats.cell(row=idx, column=1, value=name)
        stats.cell(row=idx, column=2, value=count)
    stats.column_dimensions["A"].width = 28

    ensure_output_dir()
    wb.save(OUT)
    print(f"已导出 {len(all_rows)} 条 -> {OUT}")
    return OUT


if __name__ == "__main__":
    export()
